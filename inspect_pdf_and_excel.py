import fitz
import openpyxl
import pandas as pd
import json

pdf_path = "250371509001.pdf"
excel_path = "I SEM DIGIVAL PURPOSE-25 AB F+RV.xlsx"

print("--- PDF ANALYSIS ---")
doc = fitz.open(pdf_path)
page = doc[0]
print(f"Page rect: {page.rect}")

blocks = page.get_text("dict")["blocks"]
spans_info = []
for b in blocks:
    if "lines" in b:
        for l in b["lines"]:
            for s in l["spans"]:
                spans_info.append({
                    "text": s["text"],
                    "bbox": [round(c, 2) for c in s["bbox"]],
                    "font": s["font"],
                    "size": round(s["size"], 2),
                    "flags": s["flags"],
                    "color": s["color"]
                })

with open("pdf_spans.json", "w", encoding="utf-8") as f:
    json.dump(spans_info, f, indent=2)

print(f"Extracted {len(spans_info)} text spans to pdf_spans.json.")

print("\n--- EXCEL ANALYSIS ---")
wb = openpyxl.load_workbook(excel_path, data_only=True)
print("Sheet Names:", wb.sheetnames)
for sheet in wb.sheetnames:
    df = pd.read_excel(excel_path, sheet_name=sheet)
    print(f"\nSheet: '{sheet}' - Rows: {len(df)}, Columns ({len(df.columns)}):")
    print(list(df.columns))
    if not df.empty:
        print("Sample Row 0:")
        print(df.iloc[0].to_dict())
