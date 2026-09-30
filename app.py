import os
import io
import shutil
import openpyxl
import pandas as pd
from flask import Flask, render_template, request, jsonify, send_file, url_for
import pypdfium2 as pdfium
import generator

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = os.path.join(os.getcwd(), 'static', 'uploads')
app.config['GENERATED_FOLDER'] = os.path.join(os.getcwd(), 'generated_pdfs')
app.config['SIGNATURE_FOLDER'] = os.path.join(os.getcwd(), 'static', 'uploads', 'signatures')

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['GENERATED_FOLDER'], exist_ok=True)
os.makedirs(app.config['SIGNATURE_FOLDER'], exist_ok=True)

# Default static master PDF template
DEFAULT_TEMPLATE = os.path.join(os.getcwd(), '250371509001.pdf')

generation_status = {
    "status": "idle",
    "total": 0,
    "current": 0,
    "generated": 0,
    "failed": 0,
    "message": "Ready to generate",
    "zip_available": False,
    "preview_img": None
}

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/get_sheets', methods=['POST'])
def get_sheets():
    excel_file = request.files.get('excel')
    if not excel_file or not excel_file.filename:
        return jsonify({'error': 'Please upload an Excel data file.'}), 400
        
    save_path = os.path.join(app.config['UPLOAD_FOLDER'], 'uploaded_excel.xlsx')
    excel_file.save(save_path)
        
    try:
        wb = openpyxl.load_workbook(save_path, read_only=True)
        sheets_info = []
        for s in wb.sheetnames:
            title = generator.resolve_header_title(s)
            sheets_info.append({'name': s, 'header_title': title})
        return jsonify({
            'sheets': wb.sheetnames,
            'sheets_info': sheets_info,
            'all_titles': generator.ALL_HEADER_TITLES
        })
    except Exception as e:
        return jsonify({'error': f'Failed to read Excel file: {str(e)}'}), 500

@app.route('/api/preview_signature')
def preview_signature():
    campus_id = request.args.get('campus_id', '').strip()
    if not campus_id:
        return jsonify({'error': 'Campus ID required'}), 400
    sig_folder = app.config['SIGNATURE_FOLDER']
    sig_path = generator.fetch_campus_signature(campus_id, cache_dir=sig_folder)
    if sig_path and os.path.exists(sig_path):
        return send_file(sig_path, mimetype='image/png')
    return jsonify({'error': f'Signature not found for campus {campus_id}'}), 404

@app.route('/api/generate', methods=['POST'])
def generate():
    global generation_status
    generation_status = {
        "status": "running",
        "total": 0,
        "current": 0,
        "generated": 0,
        "failed": 0,
        "message": "Initializing...",
        "zip_available": False,
        "preview_img": None
    }

    # Static Master Template from folder
    template_path = DEFAULT_TEMPLATE
    if not os.path.exists(template_path):
        fallback = os.path.join(os.getcwd(), 'pristine_template.pdf')
        if os.path.exists(fallback):
            template_path = fallback
        else:
            return jsonify({'error': 'Static master PDF template (250371509001.pdf) not found in folder.'}), 400

    # User uploaded Excel File
    excel_file = request.files.get('excel')
    if excel_file and excel_file.filename:
        excel_path = os.path.join(app.config['UPLOAD_FOLDER'], 'data.xlsx')
        excel_file.save(excel_path)
    else:
        uploaded_cache = os.path.join(app.config['UPLOAD_FOLDER'], 'uploaded_excel.xlsx')
        if os.path.exists(uploaded_cache):
            excel_path = uploaded_cache
        else:
            return jsonify({'error': 'Please upload an Excel data file before generating.'}), 400

    # Campus ID for API principal signature
    campus_id = request.form.get('campus_id', '31').strip() or '31'

    sig_folder = app.config['SIGNATURE_FOLDER']

    selected_sheet = request.form.get('sheet', '').strip()
    if not selected_sheet or selected_sheet.upper() == 'ALL':
        return jsonify({'error': 'Please select a specific sheet. Batch generation of all sheets is disabled — only one sheet can be generated at a time.'}), 400
    sheets_param = [selected_sheet]

    # Header parameters
    header_title = request.form.get('header_title', 'AUTO').strip() or 'AUTO'
    semester = request.form.get('semester', 'FIRST SEMESTER').strip() or 'FIRST SEMESTER'
    exam_month = request.form.get('exam_month', 'JANUARY').strip() or 'JANUARY'
    exam_year = request.form.get('exam_year', '2026').strip() or '2026'

    # Published Result Date (appears at bottom-left DATE : DD.MM.YYYY)
    raw_date = request.form.get('published_date') or request.form.get('issue_date') or request.form.get('fixed_date') or '2026-04-06'
    fixed_date = generator.format_display_date(raw_date)

    def update_progress(current, total, gen, fail, msg):
        global generation_status
        generation_status["current"] = current
        generation_status["total"] = total
        generation_status["generated"] = gen
        generation_status["failed"] = fail
        generation_status["message"] = msg

    try:
        res = generator.process_excel_and_generate_all(
            template_path=template_path,
            excel_path=excel_path,
            signature_folder=sig_folder,
            output_base_dir=app.config['GENERATED_FOLDER'],
            selected_sheets=sheets_param,
            date_mode="fixed",
            fixed_date=fixed_date,
            default_campus_id=campus_id,
            header_title=header_title,
            semester=semester,
            exam_month=exam_month,
            exam_year=exam_year,
            progress_callback=update_progress
        )

        preview_rel_url = None

        generation_status = {
            "status": "completed",
            "total": res["total"],
            "current": res["total"],
            "generated": res["generated"],
            "failed": res["failed"],
            "message": "Generation completed successfully!",
            "zip_available": True,
            "preview_img": preview_rel_url
        }

        return jsonify({
            "status": "completed",
            "results": res,
            "preview_img": preview_rel_url
        })

    except Exception as e:
        generation_status["status"] = "error"
        generation_status["message"] = f"Error: {str(e)}"
        return jsonify({"error": str(e)}), 500

@app.route('/api/status')
def get_status():
    return jsonify(generation_status)

@app.route('/download/zip')
def download_zip():
    zip_path = os.path.join(app.config['GENERATED_FOLDER'], 'output.zip')
    if not os.path.exists(zip_path):
        return "ZIP file not generated yet.", 404

    # Read zip file completely into memory
    with open(zip_path, 'rb') as f:
        zip_data = io.BytesIO(f.read())

    # Immediately delete all generated files from the folder
    try:
        shutil.rmtree(app.config['GENERATED_FOLDER'], ignore_errors=True)
        os.makedirs(app.config['GENERATED_FOLDER'], exist_ok=True)
    except Exception as e:
        print("Generated files cleanup error:", e)

    # Clean temporary uploaded Excel files so nothing remains stored
    try:
        excel_tmp = os.path.join(app.config['UPLOAD_FOLDER'], 'uploaded_excel.xlsx')
        if os.path.exists(excel_tmp):
            os.remove(excel_tmp)
        excel_data = os.path.join(app.config['UPLOAD_FOLDER'], 'data.xlsx')
        if os.path.exists(excel_data):
            os.remove(excel_data)
    except Exception as e:
        print("Upload cleanup error:", e)

    zip_data.seek(0)
    return send_file(
        zip_data,
        mimetype='application/zip',
        as_attachment=True,
        download_name='output.zip'
    )

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
