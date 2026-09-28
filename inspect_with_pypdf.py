import pypdf
import pandas as pd
import openpyxl
import json

print("=== EXAMINING PDF USING PYPDF ===")
reader = pypdf.PdfReader("250371509001.pdf")
print("Num pages:", len(reader.pages))
page = reader.pages[0]

def visitor_body(text, cm, tm, font_dict, font_size):
    if text.strip():
        # tm is transformation matrix: tm[4] is x, tm[5] is y
        print(f"X={tm[4]:6.1f}, Y={tm[5]:6.1f} | Size={font_size:4.1f} | Text={text.strip()!r}")

print("\n--- Extracted Text Elements with Coordinates (visitor_body) ---")
page.extract_text(visitor_text=visitor_body)

print("\n=== EXAMINING EXCEL FILE ===")
excel_file = "I SEM DIGIVAL PURPOSE-25 AB F+RV.xlsx"
wb = openpyxl.load_workbook(excel_file, data_only=True)
print("Workbook sheets:", wb.sheetnames)

for sname in wb.sheetnames:
    print(f"\n--- Sheet: '{sname}' ---")
    df = pd.read_excel(excel_file, sheet_name=sname)
    print("Columns:", list(df.columns))
    print(f"Shape: {df.shape}")
    if not df.empty:
        print("First row data:")
        for col in df.columns:
            print(f"  {col}: {df.iloc[0][col]}")
