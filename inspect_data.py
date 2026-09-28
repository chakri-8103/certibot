import fitz
import openpyxl
import pandas as pd
import json
import os

pdf_path = "250371509001.pdf"
excel_path = "I SEM DIGIVAL PURPOSE-25 AB F+RV.xlsx"

print("=== PDF INSPECTION ===")
doc = fitz.open(pdf_path)
print(f"Total pages: {len(doc)}")
page = doc[0]
rect = page.rect
print(f"Page dimensions: width={rect.width}, height={rect.height}")

text_blocks = []
blocks = page.get_text("dict")["blocks"]
for b in blocks:
    if "lines" in b:
        for l in b["lines"]:
            for s in l["spans"]:
                text_blocks.append({
                    "text": s["text"],
                    "bbox": s["bbox"],
                    "font": s["font"],
                    "size": s["size"],
                    "color": s["color"],
                    "flags": s["flags"]
                })

print(f"Found {len(text_blocks)} text spans.")
for i, span in enumerate(text_blocks):
    print(f"{i:3d}: text={span['text']!r} | bbox={[round(x, 2) for x in span['bbox']]} | font={span['font']} | size={span['size']:.2f}")

print("\n=== EXCEL INSPECTION ===")
wb = openpyxl.load_workbook(excel_path, data_only=True)
print("Sheets found:", wb.sheetnames)

for sheet in wb.sheetnames:
    print(f"\n--- Sheet: {sheet} ---")
    df = pd.read_excel(excel_path, sheet_name=sheet)
    print("Columns:", list(df.columns))
    print("Shape:", df.shape)
    print("First 2 rows:")
    print(df.head(2))
